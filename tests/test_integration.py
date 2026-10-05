"""End-to-end offline tests: synthetic 'Kite' -> cache -> resample -> strategy -> engine -> Excel."""
import json
from datetime import date, datetime, timedelta

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from config.settings import IST, AppSettings, default_backtest_period
from src.analysis.compare import (MAX_VARIANTS, Variant, compare_variants, cost_sensitivity,
                                  sensitivity_variants, split_period)
from src.data.excel_cache import ExcelCandleCache
from src.data.manager import DataManager
from src.engine.runner import BacktestRunner, DataIncompleteError
from src.reporting.comparison_report import write_comparison_report
from src.reporting.excel_report import unique_path, write_run_report
from src.strategies.registry import get_strategy
from tests.helpers import ts

INST = "NSE:SYN"
NOW = datetime(2026, 3, 1, 12, 0, tzinfo=IST)


class SynthSource:
    """Deterministic fake Kite: oscillating prices with volume bursts. Minute and daily."""

    def __init__(self):
        self.calls = []

    def _closes(self, k, n):
        i = np.arange(n)
        return 100 + 3 * np.sin((i / n * 6.0) + k * 0.9) + 0.2 * k + 0.4 * np.sin(i * 0.31)

    def fetch_candles(self, token, start, end, interval):
        self.calls.append((start, end, interval))
        rows, d = [], start
        while d <= end:
            if d.weekday() < 5:
                k = d.toordinal()
                if interval == "minute":
                    c = self._closes(k % 11, 375)
                    vol = 1000 + 1500 * (np.sin(np.arange(375) / 11.0 + k) ** 2)
                    prev = c[0]
                    for i in range(375):
                        o = prev
                        rows.append((ts(d, 9, 15) + pd.Timedelta(minutes=i), o, max(o, c[i]) + 0.05,
                                     min(o, c[i]) - 0.05, c[i], vol[i]))
                        prev = c[i]
                else:
                    base = 100 + 8 * np.sin(k / 13.0) + 0.05 * (k % 400)
                    rows.append((pd.Timestamp(d.year, d.month, d.day, tz=IST), base, base + 1.5, base - 1.5,
                                 base + 0.3 * np.sin(k), 1000 + 900 * np.sin(k / 3.0) ** 2))
            d += timedelta(days=1)
        return pd.DataFrame(rows, columns=["timestamp", "open", "high", "low", "close", "volume"])


@pytest.fixture(scope="module")
def env(tmp_path_factory):
    tmp = tmp_path_factory.mktemp("cache")
    settings = AppSettings(start_date="2026-01-12", end_date="2026-01-30", instruments=[INST])
    cache = ExcelCandleCache(tmp)
    dm = DataManager(SynthSource(), cache, settings, now_fn=lambda: NOW)
    dm.ensure_data(INST, 1, "minute", date(2025, 12, 22), date(2026, 1, 30))
    dm.ensure_data(INST, 1, "day", date(2025, 7, 1), date(2026, 1, 30))
    reports = tmp_path_factory.mktemp("reports")
    runner = BacktestRunner(settings, dm, reports)
    return settings, dm, runner, reports


def test_default_period_is_previous_six_calendar_months():
    assert default_backtest_period(date(2026, 10, 4)) == (date(2026, 4, 1), date(2026, 9, 30))
    assert default_backtest_period(date(2026, 3, 1)) == (date(2025, 9, 1), date(2026, 2, 28))
    assert default_backtest_period(date(2026, 1, 15)) == (date(2025, 7, 1), date(2025, 12, 31))


def test_settings_validation_and_roundtrip(tmp_path):
    s = AppSettings(instruments=["NSE:RELIANCE"])
    p = tmp_path / "s.json"
    s.save(p)
    assert AppSettings.load(p) == s
    assert AppSettings.load(tmp_path / "missing.json").intraday_minutes == 5
    with pytest.raises(ValueError):
        AppSettings(instruments=["RELIANCE"]).validate()
    with pytest.raises(ValueError):
        AppSettings(intraday_minutes=7).validate()
    with pytest.raises(ValueError):
        AppSettings(start_date="2026-02-01", end_date="2026-01-01").validate()
    with pytest.raises(ValueError, match="Unknown setting"):
        AppSettings.from_dict({"bogus": 1})
    p.write_text("{not json")
    with pytest.raises(ValueError, match="not valid JSON"):
        AppSettings.load(p)
    partial = AppSettings.from_dict({"intraday_costs": {"brokerage_cap": 10.0}})
    assert partial.intraday_costs["brokerage_cap"] == 10.0 and partial.intraday_costs["stt_sell_pct"] == 0.025


@pytest.mark.parametrize("key", ["ema_crossover", "orb", "swing_breakout", "swing_sr"])
def test_every_strategy_runs_through_shared_engine_and_writes_report(env, key):
    settings, dm, runner, reports = env
    overrides = {"ema_crossover": {"entry_start": "09:15"}, "orb": {"rvol_filter": False, "vwap_filter": False,
                                                                      "min_range_width_pct": 0.01,
                                                                      "max_range_width_pct": 9.0},
                 "swing_breakout": {"rvol_min": 0.5, "breakout_buffer_pct": 0.0, "max_extension_pct": 20,
                                    "breakout_lookback": 10},
                 "swing_sr": {"rvol_min": 0.5, "zone_width_atr": 0.8}}[key]
    strat = get_strategy(key, overrides)
    prepared, run = runner.run(strat, [INST], date(2026, 1, 12), date(2026, 1, 30))
    res = run.result
    # accounting identities
    if len(res.trades):
        assert abs(res.equity.equity.iloc[-1] - (settings.initial_capital + res.trades.net_pnl.sum())) < 1e-6
        assert (res.trades.costs > 0).all()
        assert ((res.trades.entry_price > 0) & (res.trades.quantity >= 1)).all()
        assert (res.trades.entry_time > res.trades.signal_time).all()           # next-candle execution
        if strat.config.intraday:
            assert (res.trades.entry_time.dt.date == res.trades.exit_time.dt.date).all()   # no overnight
            assert (res.trades.exit_time.dt.time <= pd.Timestamp("15:25").time()).all()
    path = runner.write_report(strat, prepared, run)
    assert path.exists() and path.parent == reports
    wb = load_workbook(path)
    assert wb.sheetnames[:6] == ["Summary", "Trades", "Equity Curve", "Monthly Results", "Configuration",
                                 "Data Validation"]
    ws = wb["Summary"]
    labels = {ws.cell(r, 1).value for r in range(1, 40)}
    for need in ("Initial capital", "Final capital", "Net P&L (after costs)", "Total trades", "Wins", "Losses",
                 "Win rate %", "Profit factor", "Expectancy per trade", "Maximum drawdown"):
        assert need in labels
    conf = wb["Configuration"]
    text = " ".join(str(c.value) for row in conf.iter_rows() for c in row)
    assert "risk_per_trade_pct" in text and "slippage_pct" in text and "stop is assumed to be hit first" in text


def test_intraday_strategies_do_trade_on_the_fixture(env):
    _, _, runner, _ = env
    strat = get_strategy("ema_crossover", {"entry_start": "09:15"})
    _, run = runner.run(strat, [INST], date(2026, 1, 12), date(2026, 1, 30))
    assert run.metrics["total_trades"] > 0          # makes the identities above non-vacuous


def test_report_marks_unavailable_metrics_as_na_not_zero(env, tmp_path):
    _, _, runner, _ = env
    strat = get_strategy("swing_breakout", {"rvol_min": 50})                  # impossible filter -> no trades
    prepared, run = runner.run(strat, [INST], date(2026, 1, 12), date(2026, 1, 30))
    assert run.metrics["total_trades"] == 0 and run.metrics["win_rate_pct"] is None
    path = write_run_report(run.result, runner.build_meta(strat, prepared, run.start, run.end), tmp_path)
    ws = load_workbook(path)["Summary"]
    vals = {ws.cell(r, 1).value: ws.cell(r, 2).value for r in range(1, 40)}
    assert vals["Win rate %"] == "N/A" and vals["Profit factor"] == "N/A" and vals["Expectancy per trade"] == "N/A"
    assert vals["Total trades"] == 0 and vals["Wins"] == 0                      # genuine zeros stay zeros
    assert vals["Net P&L (after costs)"] == 0


def test_reports_never_overwrite(tmp_path):
    a = unique_path(tmp_path, "run")
    a.write_text("x")
    b = unique_path(tmp_path, "run")
    assert a != b and b.name == "run_2.xlsx"


def test_backtest_refuses_incomplete_cache_but_can_be_overridden(env):
    _, _, runner, _ = env
    strat = get_strategy("swing_breakout")
    with pytest.raises(DataIncompleteError, match="Download"):
        runner.prepare(strat, [INST], date(2026, 1, 12), date(2026, 4, 30))
    prepared = runner.prepare(strat, [INST], date(2026, 1, 12), date(2026, 4, 30), require_complete=False)
    assert any("uncovered" in w for w in prepared.warnings)
    with pytest.raises(DataIncompleteError, match="no candles"):
        runner.prepare(strat, ["NSE:OTHER"], date(2026, 1, 12), date(2026, 1, 30), require_complete=False)


def test_data_validation_sheet_reports_cache_details(env, tmp_path):
    _, _, runner, _ = env
    strat = get_strategy("orb")
    prepared, run = runner.run(strat, [INST], date(2026, 1, 12), date(2026, 1, 30))
    path = runner.write_report(strat, prepared, run)
    ws = load_workbook(path)["Data Validation"]
    header = [c.value for c in ws[1]]
    row = [c.value for c in ws[2]]
    d = dict(zip(header, row))
    assert d["Source file"] == "NSE_SYN_minute.xlsx" and d["Missing ranges"] == "none"
    assert "Cache coverage" in d and d["Candles used"] > 0


# ------------------------------------------------------------------ comparison & robustness
def test_comparison_with_in_out_of_sample_and_warnings(env, tmp_path):
    _, _, runner, _ = env
    variants = [Variant("EMA", "ema_crossover", {"entry_start": "09:15"}),
                Variant("ORB", "orb", {"rvol_filter": False, "vwap_filter": False})]
    comp = compare_variants(runner, variants, [INST], date(2026, 1, 12), date(2026, 1, 30))
    assert list(comp.table.variant) == ["EMA", "ORB"]
    assert {"IS_trades", "OOS_trades", "OOS_return_pct"} <= set(comp.table.columns)
    assert comp.split_date == split_period(date(2026, 1, 12), date(2026, 1, 30), 0.7)
    assert any("trades" in w and "unstable" in w for w in comp.warnings)          # few trades warning
    assert any("NOT expected to be the best" in w for w in comp.warnings)
    path = write_comparison_report(comp, "test", [INST], tmp_path)
    wb = load_workbook(path)
    assert {"Comparison", "Monthly P&L", "Variant Config", "Warnings"} <= set(wb.sheetnames)


def test_sensitivity_is_bounded_and_skips_invalid_combinations():
    variants, notes = sensitivity_variants("ema_crossover")
    assert variants[0].label == "base" and 1 < len(variants) <= MAX_VARIANTS
    assert any("fast_period" in n for n in notes) or all(v.overrides is not None for v in variants)
    with pytest.raises(ValueError, match="Grid would create"):
        sensitivity_variants("ema_crossover", grid={"fast_period": list(range(2, 10)),
                                                    "slow_period": list(range(11, 20)),
                                                    "reward_to_risk": [1, 2, 3]}, full_grid=True)


def test_cost_sensitivity_more_costs_never_improve_pnl(env):
    _, _, runner, _ = env
    table = cost_sensitivity(runner, "ema_crossover", [INST], date(2026, 1, 12), date(2026, 1, 30),
                             overrides={"entry_start": "09:15"})
    pnl = table.net_pnl.tolist()
    assert table.cost_multiplier.tolist() == [0.0, 1.0, 2.0, 3.0]
    assert pnl == sorted(pnl, reverse=True) and table.total_costs.iloc[0] == 0
