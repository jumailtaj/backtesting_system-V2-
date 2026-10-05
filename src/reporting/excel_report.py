"""Per-run Excel report.  Files are never overwritten (a numeric suffix is added)."""
from __future__ import annotations

import re
from datetime import date, datetime
from pathlib import Path
from typing import Optional

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from config.settings import IST, REPORT_DIR
from src.reporting.metrics import compute_metrics, monthly_results, stability_warnings

NA = "N/A"
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=10)
BODY_FONT = Font(name="Arial", size=10)
NA_FONT = Font(name="Arial", size=10, italic=True, color="808080")
TITLE_FONT = Font(name="Arial", bold=True, size=14)
THIN = Side(style="thin", color="BFBFBF")
MONEY, PCT, NUM, TS = "#,##0.00;[Red]-#,##0.00", "0.00", "#,##0.00", "yyyy-mm-dd hh:mm"


def unique_path(directory: Path, stem: str, suffix: str = ".xlsx") -> Path:
    directory.mkdir(parents=True, exist_ok=True)
    stem = re.sub(r"[^A-Za-z0-9_.\-]", "_", stem)
    path = directory / f"{stem}{suffix}"
    n = 2
    while path.exists():
        path = directory / f"{stem}_{n}{suffix}"
        n += 1
    return path


def _naive(v):
    if isinstance(v, pd.Timestamp) and v.tzinfo is not None:
        return v.tz_convert(IST).tz_localize(None).to_pydatetime()
    if isinstance(v, pd.Timestamp):
        return v.to_pydatetime()
    if isinstance(v, pd.Timedelta):
        return str(v)
    return v


def write_table(ws, df: pd.DataFrame, formats: dict, start_row: int = 1, widths: Optional[dict] = None) -> None:
    for j, col in enumerate(df.columns, 1):
        c = ws.cell(row=start_row, column=j, value=col)
        c.font, c.fill = HEADER_FONT, HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, row in enumerate(df.itertuples(index=False), start_row + 1):
        for j, (col, v) in enumerate(zip(df.columns, row), 1):
            v = _naive(v)
            if v is None or (isinstance(v, float) and pd.isna(v)) or v is pd.NaT:
                cell = ws.cell(row=i, column=j, value=NA)
                cell.font = NA_FONT
                continue
            if hasattr(v, "item"):
                v = v.item()
            cell = ws.cell(row=i, column=j, value=v)
            cell.font = BODY_FONT
            if col in formats:
                cell.number_format = formats[col]
    for j, col in enumerate(df.columns, 1):
        w = (widths or {}).get(col)
        if w is None:
            sample = [len(str(col))] + [len(str(_naive(v))) for v in df[col].head(200)]
            w = min(max(sample) + 2, 48)
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)


def _kv_sheet(ws, rows: list, title: str) -> None:
    ws["A1"], ws["A1"].font = title, TITLE_FONT
    for i, (k, v, fmt) in enumerate(rows, 3):
        a, b = ws.cell(row=i, column=1, value=k), ws.cell(row=i, column=2)
        a.font = Font(name="Arial", bold=True, size=10)
        if v is None or (isinstance(v, float) and pd.isna(v)):
            b.value, b.font = NA, NA_FONT
        else:
            b.value, b.font = _naive(v), BODY_FONT
            if fmt:
                b.number_format = fmt
        b.alignment = Alignment(horizontal="left")
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 60


def _flatten_config(cfg: dict, prefix: str = "") -> list:
    rows = []
    for k, v in cfg.items():
        key = f"{prefix}{k}"
        if isinstance(v, dict):
            rows += _flatten_config(v, key + ".")
        else:
            rows.append((key, str(v) if isinstance(v, (list, tuple)) else v, None))
    return rows


def write_run_report(result, meta: dict, directory: Path = REPORT_DIR) -> Path:
    """``meta`` keys: strategy_label, strategy_key, instruments, timeframe, start, end, config (dict of
    sections), data_info (list of per-instrument dicts), assumptions (list[str]), extra_warnings (list)."""
    m = compute_metrics(result.trades, result.equity, result.initial_capital)
    monthly = monthly_results(result.trades, result.equity, result.initial_capital, meta["start"], meta["end"])
    warnings = stability_warnings(m, meta.get("min_trades_warning", 30)) + list(result.warnings) + \
        list(meta.get("extra_warnings", []))
    wb = Workbook()

    # ---- Summary
    ws = wb.active
    ws.title = "Summary"
    rows = [
        ("Strategy", meta["strategy_label"], None), ("Instrument(s)", ", ".join(meta["instruments"]), None),
        ("Timeframe", meta["timeframe"], None), ("Period", f"{meta['start']} to {meta['end']}", None),
        ("Report generated (IST)", datetime.now(IST).strftime("%Y-%m-%d %H:%M:%S"), None),
        ("Initial capital", m["initial_capital"], MONEY), ("Final capital", m["final_capital"], MONEY),
        ("Net P&L (after costs)", m["net_pnl"], MONEY), ("Return %", m["return_pct"], PCT),
        ("Total trades", m["total_trades"], "0"), ("Wins", m["wins"], "0"), ("Losses", m["losses"], "0"),
        ("Break-even trades", m["breakeven"], "0"), ("Win rate %", m["win_rate_pct"], PCT),
        ("Profit factor", m["profit_factor"], NUM), ("Expectancy per trade", m["expectancy"], MONEY),
        ("Average R per trade", m["avg_r"], NUM), ("Average win", m["avg_win"], MONEY),
        ("Average loss", m["avg_loss"], MONEY), ("Gross P&L", m["gross_pnl"], MONEY),
        ("Total costs", m["total_costs"], MONEY), ("Maximum drawdown", m["max_drawdown"], MONEY),
        ("Maximum drawdown %", m["max_drawdown_pct"], PCT),
        ("Max consecutive losses", m["max_consecutive_losses"], "0"),
        ("Signals generated", result.signals_generated, "0"), ("Signals rejected", len(result.rejections), "0"),
    ]
    _kv_sheet(ws, rows, f"{meta['strategy_label']} - Backtest Summary")
    r = len(rows) + 4
    ws.cell(row=r, column=1, value="Notes and warnings").font = Font(name="Arial", bold=True, size=10)
    for k, w in enumerate(warnings or ["none"], r + 1):
        ws.cell(row=k, column=1, value=w).font = BODY_FONT
    k = r + len(warnings or [1]) + 2
    ws.cell(row=k, column=1, value="N/A means the metric is unavailable or undefined (e.g. profit factor with no "
                                   "losing trades) -- it is NOT zero.").font = NA_FONT
    ws.cell(row=k + 1, column=1, value="Historical results are not a guarantee of future performance.").font = NA_FONT

    # ---- Trades
    ws = wb.create_sheet("Trades")
    t = result.trades.copy()
    if len(t):
        t["holding_time"] = t["holding_time"].astype(str)
    t = t.rename(columns={
        "instrument": "Instrument", "strategy": "Strategy", "direction": "Direction", "signal_time": "Signal time",
        "entry_time": "Entry time", "exit_time": "Exit time", "quantity": "Quantity", "entry_price": "Entry price",
        "exit_price": "Exit price", "stop_price": "Stop", "target_price": "Target", "gross_pnl": "Gross P&L",
        "costs": "Costs", "net_pnl": "Net P&L", "r_multiple": "Return (R)", "holding_time": "Holding time",
        "exit_reason": "Exit reason", "setup": "Setup",
        "stop_and_target_same_candle": "Stop & target same candle (stop assumed first)"})
    write_table(ws, t, {"Signal time": TS, "Entry time": TS, "Exit time": TS, "Entry price": NUM, "Exit price": NUM,
                        "Stop": NUM, "Target": NUM, "Gross P&L": MONEY, "Costs": MONEY, "Net P&L": MONEY,
                        "Return (R)": NUM, "Quantity": "0"})

    # ---- Equity curve
    ws = wb.create_sheet("Equity Curve")
    e = result.equity.rename(columns={"timestamp": "Timestamp (candle start)", "equity": "Equity",
                                      "drawdown": "Drawdown", "drawdown_pct": "Drawdown %"})
    write_table(ws, e, {"Timestamp (candle start)": TS, "Equity": MONEY, "Drawdown": MONEY, "Drawdown %": PCT})
    if len(e) > 1:
        chart = LineChart()
        chart.title, chart.height, chart.width = "Equity", 8, 18
        chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=min(len(e) + 1, 20000)), titles_from_data=True)
        chart.legend = None
        ws.add_chart(chart, "G2")

    # ---- Monthly
    ws = wb.create_sheet("Monthly Results")
    write_table(ws, monthly.rename(columns={"month": "Month", "trades": "Trades", "wins": "Wins",
                                            "net_pnl": "Net P&L", "return_pct": "Return %",
                                            "end_equity": "End equity"}),
                {"Net P&L": MONEY, "Return %": PCT, "End equity": MONEY})

    # ---- Configuration
    ws = wb.create_sheet("Configuration")
    rows = []
    for section, content in meta["config"].items():
        rows.append((f"== {section} ==", None, None))
        rows += _flatten_config(content)
    rows.append(("== Execution assumptions ==", None, None))
    rows += [(f"{i}", a, None) for i, a in enumerate(meta.get("assumptions", []), 1)]
    _kv_sheet(ws, rows, "Configuration")

    # ---- Data validation
    ws = wb.create_sheet("Data Validation")
    di = []
    for d in meta["data_info"]:
        di.append({
            "Instrument": d["instrument"], "Timeframe": d["timeframe"], "Source file": d["source_file"],
            "Requested (incl. warm-up)": f"{d['requested'][0]} to {d['requested'][1]}",
            "Cache coverage": "; ".join(f"{a} to {b}" for a, b in d["covered_ranges"]) or "none",
            "Missing ranges": "; ".join(f"{a} to {b}" for a, b in d["missing_ranges"]) or "none",
            "Candles used": d.get("rows"), "Raw rows read": d.get("raw_rows"),
            "Weekdays with no data (likely holidays)": len(d.get("no_data_weekdays", [])),
            "Partial sessions (<375 minute rows)": len(d.get("partial_sessions", {})),
            "Duplicate rows ignored": d.get("duplicates_removed", 0),
            "Invalid rows ignored": str(d.get("invalid_rows") or {}) if d.get("invalid_rows") else "0",
            "Incomplete candles dropped": (d.get("resample") or {}).get("incomplete_dropped", 0),
            "Warnings": " | ".join(d.get("warnings", [])) or "none",
        })
    write_table(ws, pd.DataFrame(di), {}, widths={"Warnings": 70, "Cache coverage": 40, "Missing ranges": 30})
    nodata = [(d["instrument"], x) for d in meta["data_info"] for x in d.get("no_data_weekdays", [])][:100]
    if nodata:
        r0 = len(di) + 4
        ws.cell(row=r0, column=1, value="Weekdays inside covered ranges with no candles (first 100)").font = \
            Font(name="Arial", bold=True, size=10)
        for k, (inst, day) in enumerate(nodata, r0 + 1):
            ws.cell(row=k, column=1, value=inst).font = BODY_FONT
            ws.cell(row=k, column=2, value=day).number_format = "yyyy-mm-dd"

    # ---- Rejections (extra, useful sheet)
    ws = wb.create_sheet("Rejected Signals")
    rj = result.rejections.rename(columns={"time": "Signal/entry time", "instrument": "Instrument",
                                           "reason": "Reason", "detail": "Detail"})
    write_table(ws, rj, {"Signal/entry time": TS})
    if len(rj):
        counts = result.rejections["reason"].value_counts().reset_index()
        counts.columns = ["Reason", "Count"]
        ws.cell(row=1, column=7, value="Reason summary").font = Font(name="Arial", bold=True)
        for k, (reason, n) in enumerate(counts.itertuples(index=False), 2):
            ws.cell(row=k, column=7, value=reason).font = BODY_FONT
            ws.cell(row=k, column=8, value=int(n)).font = BODY_FONT
        ws.column_dimensions["G"].width = 38

    stem = (f"{meta['strategy_key']}_{'-'.join(i.replace(':', '-') for i in meta['instruments'][:3])}_"
            f"{meta['timeframe']}_{meta['start']}_{meta['end']}_{datetime.now(IST):%Y%m%d-%H%M%S}")
    path = unique_path(directory, stem)
    wb.save(path)
    return path
