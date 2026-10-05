"""Excel report for strategy/variant comparisons."""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook

from config.settings import IST, REPORT_DIR
from src.reporting.excel_report import MONEY, NUM, PCT, unique_path, write_table


def write_comparison_report(comp, label: str, instruments: list, directory: Path = REPORT_DIR) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Comparison"
    t = comp.table.copy()
    fmts = {c: PCT for c in t.columns if c.endswith("_pct")}
    fmts.update({c: NUM for c in t.columns if "profit_factor" in c or c == "expectancy"})
    fmts.update({"net_pnl": MONEY, "total_costs": MONEY})
    write_table(ws, t, fmts)
    r = len(t) + 3
    ws.cell(row=r, column=1, value=f"Period {comp.start} to {comp.end}" + (
        f"; in-sample to {comp.split_date}, out-of-sample after" if comp.split_date else ""))
    ws.cell(row=r + 1, column=1, value="N/A = metric unavailable or undefined (not zero).")

    if comp.cost_table is not None:
        write_table(wb.create_sheet("Cost Sensitivity"), comp.cost_table,
                    {"net_pnl": MONEY, "return_pct": PCT, "profit_factor": NUM, "expectancy": NUM,
                     "win_rate_pct": PCT, "max_drawdown_pct": PCT, "total_costs": MONEY})
    ws = wb.create_sheet("Monthly P&L")
    if len(comp.monthly):
        pivot = comp.monthly.pivot_table(index="variant", columns="month", values="net_pnl", aggfunc="sum",
                                         sort=False).reset_index()
        pivot.columns = [str(c) for c in pivot.columns]
        write_table(ws, pivot, {c: MONEY for c in pivot.columns if c != "variant"})
    ws = wb.create_sheet("Variant Config")
    cfg = pd.DataFrame(comp.configs).reset_index().rename(columns={"index": "parameter"})
    write_table(ws, cfg.astype(object).where(cfg.notna(), None), {})
    ws = wb.create_sheet("Warnings")
    for i, w in enumerate(comp.warnings or ["none"], 1):
        ws.cell(row=i, column=1, value=w)
    ws.column_dimensions["A"].width = 140
    stem = f"comparison_{label}_{'-'.join(i.replace(':', '-') for i in instruments[:3])}_" \
           f"{comp.start}_{comp.end}_{datetime.now(IST):%Y%m%d-%H%M%S}"
    path = unique_path(directory, stem)
    wb.save(path)
    return path
